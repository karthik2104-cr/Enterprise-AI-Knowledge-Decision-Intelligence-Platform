from pathlib import Path

from fastapi import HTTPException

SUPPORTED = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}


def extract_text(path: Path) -> list[tuple[int | None, str, str | None]]:
    """Return list of (page, text, section)."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {suffix}")
    try:
        if suffix == ".pdf":
            return _extract_pdf(path)
        if suffix == ".docx":
            return _extract_docx(path)
        text = path.read_text(encoding="utf-8", errors="replace")
        if not text.strip():
            raise HTTPException(status_code=400, detail="The uploaded file is empty.")
        return [(1, text, _first_heading(text))]
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Failed to extract text: {exc}") from exc


def _first_heading(text: str) -> str | None:
    for line in text.splitlines():
        stripped = line.strip().lstrip("#").strip()
        if stripped:
            return stripped[:120]
    return None


def _extract_pdf(path: Path) -> list[tuple[int | None, str, str | None]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[tuple[int | None, str, str | None]] = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append((i, text, _first_heading(text)))
    joined = "\n".join(p[1] for p in pages).strip()
    if not joined:
        raise HTTPException(status_code=400, detail="PDF contained no extractable text.")
    return pages


def _extract_docx(path: Path) -> list[tuple[int | None, str, str | None]]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    text = "\n".join(paragraphs)
    if not text.strip():
        raise HTTPException(status_code=400, detail="DOCX contained no extractable text.")
    heading = None
    for p in doc.paragraphs:
        if p.style and p.style.name and str(p.style.name).startswith("Heading") and p.text.strip():
            heading = p.text.strip()[:120]
            break
    return [(1, text, heading or _first_heading(text))]
